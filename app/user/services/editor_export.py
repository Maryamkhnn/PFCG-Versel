from io import BytesIO
from html import escape


def paragraphs(delta):
    runs = []

    for operation in delta["ops"]:
        attributes = operation.get("attributes", {})
        pieces = operation["insert"].split("\n")

        for index, piece in enumerate(pieces):
            if piece:
                runs.append((piece, attributes))

            if index < len(pieces) - 1:
                yield runs, attributes
                runs = []


def color(value):
    value = value.lstrip("#")

    if len(value) == 3:
        return "".join(character * 2 for character in value)

    return value


def export_word(title, delta):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()

    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)

    document.add_heading(title, 0)

    for runs, attributes in paragraphs(delta):

        style = (
            "Heading " + str(attributes["header"])
            if attributes.get("header")
            else None
        )

        if not style and attributes.get("list"):
            style = (
                "List Number"
                if attributes["list"] == "ordered"
                else "List Bullet"
            )

        paragraph = document.add_paragraph(style=style)

        paragraph.alignment = {
            "center": 1,
            "right": 2,
            "justify": 3
        }.get(attributes.get("align"), 0)

        paragraph.paragraph_format.line_spacing = float(
            attributes.get("lineheight", "1.5")
        )

        paragraph.paragraph_format.left_indent = Inches(
            attributes.get("indent", 0) * 0.2
        )

        for value, formatting in runs:
            run = paragraph.add_run(value)

            for key in ("bold", "italic", "underline"):
                if key in formatting:
                    setattr(run, key, formatting[key])

            run.font.strike = formatting.get("strike", False)

            if formatting.get("font"):
                run.font.name = {
                    "serif": "Times New Roman",
                    "monospace": "Courier New"
                }[formatting["font"]]

            if formatting.get("size"):
                run.font.size = Pt({
                    "small": 9,
                    "large": 16,
                    "huge": 24
                }[formatting["size"]])

            if formatting.get("color"):
                run.font.color.rgb = RGBColor.from_string(
                    color(formatting["color"])
                )

            if formatting.get("background"):
                shading = OxmlElement("w:shd")
                shading.set(
                    qn("w:fill"),
                    color(formatting["background"])
                )
                run._r.get_or_add_rPr().append(shading)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output


def export_pdf(title, delta):
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer
    )
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=48,
        rightMargin=48,
        topMargin=48,
        bottomMargin=48,
        title=title
    )

    story = [
        Paragraph(
            escape(title),
            ParagraphStyle(
                "title",
                fontName="Helvetica-Bold",
                fontSize=20,
                leading=26,
                spaceAfter=18,
                textColor="#1f1648"
            )
        )
    ]

    counters = {}

    for runs, attributes in paragraphs(delta):
        heading_level = attributes.get("header", 0)
        size = {1: 20, 2: 16, 3: 14}.get(heading_level, 11)

        fragments = []
        largest = size

        for value, formatting in runs:
            family = {
                "serif": "Times-Roman",
                "monospace": "Courier"
            }.get(formatting.get("font"), "Helvetica")

            font_size = {
                "small": 9,
                "large": 16,
                "huge": 24
            }.get(formatting.get("size"), size)

            largest = max(largest, font_size)

            parameters = f'name="{family}" size="{font_size}"'

            if formatting.get("color"):
                parameters += (
                    f' color="#{color(formatting["color"])}"'
                )

            if formatting.get("background"):
                parameters += (
                    f' backColor="#{color(formatting["background"])}"'
                )

            fragment = (
                f"<font {parameters}>{escape(value)}</font>"
            )

            for key, tag in (
                ("bold", "b"),
                ("italic", "i"),
                ("underline", "u"),
                ("strike", "strike")
            ):
                enabled = formatting.get(key)

                if (
                    key == "bold"
                    and heading_level
                    and formatting.get(key) is not False
                ):
                    enabled = True

                if enabled:
                    fragment = f"<{tag}>{fragment}</{tag}>"

            fragments.append(fragment)

        indent = attributes.get("indent", 0)
        prefix = ""

        if attributes.get("list") == "ordered":
            counters = {
                key: value
                for key, value in counters.items()
                if key <= indent
            }
            counters[indent] = counters.get(indent, 0) + 1
            prefix = str(counters[indent]) + ". "

        elif attributes.get("list") == "bullet":
            counters.pop(indent, None)
            prefix = "• "

        else:
            counters.clear()

        style = ParagraphStyle(
            "body",
            fontName="Helvetica",
            fontSize=size,
            leading=max(
                largest * float(attributes.get("lineheight", "1.5")),
                largest + 2
            ),
            spaceAfter=8,
            leftIndent=indent * 14,
            alignment={
                "center": 1,
                "right": 2,
                "justify": 4
            }.get(attributes.get("align"), 0),
            splitLongWords=True,
            allowWidows=1,
            allowOrphans=1
        )

        if runs:
            story.append(
                Paragraph(prefix + "".join(fragments), style)
            )
        else:
            story.append(Spacer(1, 10))

    def footer(canvas, page):
        canvas.saveState()
        canvas.setFont("Helvetica", 9)
        canvas.setFillColorRGB(0.5, 0.45, 0.55)

        canvas.drawString(48, 26, "PFCG AI")

        canvas.drawRightString(
            A4[0] - 48,
            26,
            str(page.page)
        )

        canvas.restoreState()

    document.build(
        story,
        onFirstPage=footer,
        onLaterPages=footer
    )

    output.seek(0)
    return output